"""Фикстуры: генерим tiny .docx рядом с sample.txt."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session")
def sample_txt() -> Path:
    path = FIXTURES / "sample.txt"
    assert path.is_file()
    return path


def _write_sample_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("Sample Title", level=1)
    doc.add_paragraph("Hello from the sample fixture paragraph.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "42"
    doc.add_paragraph("Closing paragraph after the table.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def _write_rich_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("Rich Sample Document", level=1)
    doc.add_paragraph(
        "Introductory paragraph under H1. This doc exercises headings, lists, and a table."
    )
    doc.add_heading("Section Two — Details", level=2)
    doc.add_paragraph("Body paragraph under H2 with enough prose for section_path checks.")
    doc.add_paragraph("First bullet item", style="List Bullet")
    doc.add_paragraph("Second bullet item", style="List Bullet")
    doc.add_paragraph("Third bullet item", style="List Bullet")
    doc.add_heading("Metrics Table", level=2)
    table = doc.add_table(rows=4, cols=3)
    rows = [
        ["Metric", "Q1", "Q2"],
        ["Revenue", "100", "120"],
        ["Users", "50", "75"],
        ["NPS", "42", "55"],
    ]
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            table.cell(r_i, c_i).text = val
    doc.add_paragraph("Closing paragraph after the metrics table.")
    doc.add_heading("Appendix", level=2)
    doc.add_paragraph("Appendix note for nested section_path.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


@pytest.fixture(scope="session")
def sample_docx(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Минимальный docx с заголовком, абзацем и таблицей."""
    out = FIXTURES / "sample.docx"
    _write_sample_docx(out)
    alt = tmp_path_factory.mktemp("docx") / "sample.docx"
    _write_sample_docx(alt)
    return out if out.is_file() else alt


@pytest.fixture(scope="session")
def rich_docx(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Богатый docx: H1/H2, список, таблица — для section_path / labels."""
    out = FIXTURES / "rich_sample.docx"
    _write_rich_docx(out)
    alt = tmp_path_factory.mktemp("rich") / "rich_sample.docx"
    _write_rich_docx(alt)
    return out if out.is_file() else alt


def _write_docx_with_image(path: Path) -> Path:
    """Мини-PNG внутри docx — проверка image branch."""
    from docx.shared import Inches
    from PIL import Image as PILImage

    png_path = path.with_suffix(".png")
    PILImage.new("RGB", (32, 24), color=(20, 120, 200)).save(png_path)

    doc = Document()
    doc.add_heading("Doc With Image", level=1)
    doc.add_paragraph("Paragraph before picture.")
    doc.add_picture(str(png_path), width=Inches(1.0))
    doc.add_paragraph("Paragraph after picture.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


@pytest.fixture(scope="session")
def docx_with_image(tmp_path_factory: pytest.TempPathFactory) -> Path:
    out = FIXTURES / "with_image.docx"
    _write_docx_with_image(out)
    alt = tmp_path_factory.mktemp("img") / "with_image.docx"
    _write_docx_with_image(alt)
    return out if out.is_file() else alt
